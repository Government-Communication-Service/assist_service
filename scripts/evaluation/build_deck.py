# /// script
# requires-python = ">=3.12"
# dependencies = ["pandas>=2.2", "matplotlib>=3.8", "python-pptx>=1.0"]
# ///
# ruff: noqa: E501, E731, E741, C408, B008, B023
"""Build the summary slide deck for one evaluation run.

Run from scripts/evaluation after judge.py collect and collate.py:
    uv run build_deck.py --run-id 2026-10-08-always_rag
Writes evaluation-summary-<run-id>.pptx next to the HTML reports. The candidate (default agent_loop) is compared with the baseline (default dev_main); name others with --candidate and --baseline.
"""

import argparse
import json
import tempfile
from pathlib import Path

import matplotlib
import pandas as pd

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

REPO = Path(__file__).resolve().parents[2]
parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
parser.add_argument("--run-id", required=True)
parser.add_argument(
    "--workspace", type=Path, default=REPO.parent / "copilot-evaluation", help="Judge workspace holding verdicts/."
)
parser.add_argument("--out-dir", type=Path, default=REPO / "data/ignored/evaluation")
parser.add_argument(
    "--examples",
    default="govuk-003:4:2;long_chat-001:2:11;govuk-002:5:1;smart_targets-001:3:1;no_tool-001:3:1;document_summary-001:3:0",
    help="Appendix example turns as case_id:repeat:turn_index, separated by semicolons.",
)
parser.add_argument("--candidate", default="agent_loop", help="Environment being evaluated (the new feature).")
parser.add_argument("--baseline", default="dev_main", help="Environment compared against (for example prod).")
parser.add_argument("--candidate-label", help="Display name for the candidate (default: the environment name).")
parser.add_argument("--baseline-label", help="Display name for the baseline (default: the environment name).")
args = parser.parse_args()
RUN = args.run_id
R = f"{args.out_dir}/results/{RUN}/"
OUT = f"{args.out_dir}/"
HERE = tempfile.mkdtemp(prefix="deck-")
WORKSPACE = args.workspace
CAND, BASE = args.candidate, args.baseline
COL = {CAND: "#4c78a8", BASE: "#f58518", "tie": "#d9d9d9"}
KNOWN = {"agent_loop": "Agent loop", "dev_main": "Dev-main", "prod": "Prod"}
CL = args.candidate_label or KNOWN.get(args.candidate, args.candidate.replace("_", " ").capitalize())
BL = args.baseline_label or KNOWN.get(args.baseline, args.baseline.replace("_", " ").capitalize())
LABEL = {CAND: CL, BASE: BL, "tie": "Tie"}

t = pd.DataFrame(json.load(open(R + "collated/turns.json")))
W = json.load(open(R + "collated/windows.json"))
J = pd.DataFrame([json.loads(l) for l in open(R + "judgments.jsonl")])
for c in ("input_tokens", "cache_read_tokens", "cache_write_tokens", "output_tokens"):
    t[c] = t[c].fillna(0)
t["total_in"] = t.input_tokens + t.cache_read_tokens + t.cache_write_tokens
t["failed"] = t.error.notna()
assert not t.failed.any()
KEYS = ["case_id", "repeat", "turn_index"]

# ---- statistics
S = {}
for e in (CAND, BASE):
    r = t[t.env == e]
    other = sum(p["cost"] for p in W[e]["by_purpose"] if p["purpose"] != "chat_loop")
    S[e] = dict(
        answer=r.answer_cost.sum(),
        other=other,
        total=r.answer_cost.sum() + other,
        cache=r.cache_read_tokens.sum() / r.total_in.sum(),
        ttft50=r.time_to_first_token_s.quantile(0.5),
        ttft95=r.time_to_first_token_s.quantile(0.95),
        tt50=r.total_time_s.quantile(0.5),
        tt95=r.total_time_s.quantile(0.95),
        other_calls=sum(p["calls"] for p in W[e]["by_purpose"] if p["purpose"] != "chat_loop"),
    )
acc = J[J.kind == "acceptable"]
for e in (CAND, BASE):
    S[e]["acc"] = acc[acc.env == e].acceptable.sum()
    S[e]["acc_n"] = int((acc.env == e).sum())
pw = J[J.kind == "pairwise"]
wins = pw.winner.value_counts().to_dict()
n_scen, n_cases, n_turn_defs = (
    t.scenario.nunique(),
    t.case_id.nunique(),
    t[["case_id", "turn_index"]].drop_duplicates().shape[0],
)
reps = t.repeat.nunique()
per_env_turns = int((t.env == CAND).sum())

# ---- charts
plt.rcParams.update(
    {"font.family": "DejaVu Sans", "font.size": 10, "axes.spines.top": False, "axes.spines.right": False}
)
scen = sorted(t.scenario.unique())


def save(fig, name):
    fig.savefig(f"{HERE}/{name}.png", dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig)


# cumulative cost, long_chat
lc = t[t.scenario == "long_chat"].sort_values(["env", "case_id", "repeat", "turn_index"]).copy()
lc["cum"] = lc.groupby(["env", "case_id", "repeat"]).answer_cost.cumsum()
sc = lc.groupby(["env", "turn_index"], as_index=False).cum.mean()
fig, ax = plt.subplots(figsize=(6.2, 3.9))
for e in (CAND, BASE):
    d = sc[sc.env == e]
    ax.plot(d.turn_index, d.cum, marker="o", ms=4, color=COL[e], label=LABEL[e])
ax.set_xlabel("Turn in the long chat")
ax.set_ylabel("Cumulative answer cost ($)")
ax.set_title("Long chat: cumulative answer cost (mean of repeats)", fontsize=11, loc="left")
ax.legend(frameon=False)
ax.grid(axis="y", alpha=0.3)
ax.yaxis.set_major_formatter(lambda v, _: f"${v:.2f}")
save(fig, "cum")


def boxes(field, title, name):
    fig, ax = plt.subplots(figsize=(6.4, 3.6))
    w = 0.36
    for j, e in enumerate((CAND, BASE)):
        data = [t[(t.env == e) & (t.scenario == s)][field].dropna().values for s in scen]
        pos = [i + (j - 0.5) * (w + 0.02) for i in range(len(scen))]
        ax.boxplot(
            data,
            positions=pos,
            widths=w,
            whis=(0, 100),
            patch_artist=True,
            showfliers=False,
            medianprops=dict(color="black"),
            boxprops=dict(facecolor=COL[e], alpha=0.85, edgecolor=COL[e]),
            whiskerprops=dict(color=COL[e]),
            capprops=dict(color=COL[e]),
        )
        ax.plot([], [], color=COL[e], lw=8, label=LABEL[e])
    ax.set_xticks(range(len(scen)))
    ax.set_xticklabels(
        [s.replace("document_", "doc-").replace("_", "\n") for s in scen], fontsize=8, rotation=35, ha="right"
    )
    ax.set_ylabel("seconds")
    ax.set_title(title, fontsize=11, loc="left")
    ax.grid(axis="y", alpha=0.3)
    ax.legend(frameon=False, fontsize=8, loc="upper left")
    save(fig, name)


boxes("time_to_first_token_s", "Time to first token (min to max, box = quartiles)", "ttft")
boxes("total_time_s", "Total time (min to max, box = quartiles)", "total")

# acceptable
m = t.merge(acc[["env"] + KEYS + ["acceptable"]], on=["env"] + KEYS)
rate = m.groupby(["scenario", "env"]).acceptable.mean().unstack()
fig, ax = plt.subplots(figsize=(6.4, 3.6))
w = 0.38
for j, e in enumerate((CAND, BASE)):
    ax.bar(
        [i + (j - 0.5) * w for i in range(len(scen))], [rate.loc[s, e] for s in scen], w, color=COL[e], label=LABEL[e]
    )
ax.set_xticks(range(len(scen)))
ax.set_xticklabels(
    [s.replace("document_", "doc-").replace("_", "\n") for s in scen], fontsize=8, rotation=35, ha="right"
)
ax.set_ylim(0, 1)
ax.yaxis.set_major_formatter(lambda v, _: f"{v:.0%}")
ax.grid(axis="y", alpha=0.3)
ax.set_title("Share of answers judged acceptable", fontsize=11, loc="left")
ax.legend(frameon=False, fontsize=8, ncol=2, loc="upper center", bbox_to_anchor=(0.5, -0.28))
save(fig, "acc")

# pairwise
sk = t[["scenario"] + KEYS].drop_duplicates()
pm = pw.merge(sk, on=KEYS).groupby(["scenario", "winner"]).size().unstack(fill_value=0)
pm = pm.div(pm.sum(axis=1), axis=0)
fig, ax = plt.subplots(figsize=(6.4, 3.6))
bottom = pd.Series(0.0, index=pm.index)
for k in (CAND, "tie", BASE):
    v = pm.get(k, pd.Series(0.0, index=pm.index)).reindex(scen)
    ax.bar(range(len(scen)), v.values, 0.6, bottom=bottom.reindex(scen).values, color=COL[k], label=LABEL[k])
    bottom = bottom.add(v, fill_value=0)
ax.set_xticks(range(len(scen)))
ax.set_xticklabels(
    [s.replace("document_", "doc-").replace("_", "\n") for s in scen], fontsize=8, rotation=35, ha="right"
)
ax.yaxis.set_major_formatter(lambda v, _: f"{v:.0%}")
ax.set_title("Pairwise winners", fontsize=11, loc="left")
ax.legend(frameon=False, fontsize=8, ncol=3, loc="upper center", bbox_to_anchor=(0.5, -0.28))
save(fig, "pair")

# ---- swap / bias stats
key = json.load(open(R + "judge_key.json"))["items"]
V = {i: json.load(open(f"{WORKSPACE}/verdicts/{i}.json")) for i in key}
pos = pd.Series([v["winner"] for v in V.values()]).value_counts().to_dict()
by = {}
for i, it in key.items():
    by.setdefault((it["case_id"], it["repeat"], it["turn_index"]), {})[it["is_swap"]] = (it, V[i])
wv = lambda it, v: {"A": it["a_env"], "B": it["b_env"], "tie": "tie"}[v["winner"]]
pairs = [p for p in by.values() if len(p) == 2]
same = sum(wv(*p[False]) == wv(*p[True]) for p in pairs)
tie_vs_win = sum(1 for p in pairs if wv(*p[False]) != wv(*p[True]) and "tie" in (wv(*p[False]), wv(*p[True])))
flips = len(pairs) - same - tie_vs_win
tw = [
    wv(*p[False]) if wv(*p[True]) == "tie" else wv(*p[True])
    for p in pairs
    if wv(*p[False]) != wv(*p[True]) and "tie" in (wv(*p[False]), wv(*p[True]))
]
n_swaps = sum(it["is_swap"] for it in key.values())

# ---- deck
NAVY = RGBColor(0x1D, 0x2B, 0x4F)
GREY = RGBColor(0x55, 0x5B, 0x66)
LIGHT = RGBColor(0xF2, 0xF4, 0xF8)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]


def tb(slide, x, y, w, h, paras, size=14, color=GREY, bold=False, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, p in enumerate(paras if isinstance(paras, list) else [paras]):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        runs = p if isinstance(p, list) else [(p, bold)]
        for text, b in runs:
            r = para.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = b
            r.font.color.rgb = color
            r.font.name = "Calibri"
        para.space_after = Pt(6)
    return box


def new(title, sub=None):
    s = prs.slides.add_slide(blank)
    tb(s, 0.6, 0.35, 12.1, 0.8, title, size=30, color=NAVY, bold=True)
    if sub:
        tb(s, 0.6, 1.05, 12.1, 0.5, sub, size=15)
    return s


def table(slide, x, y, w, rows, colw, size=13, rowh=0.42, winner_col=None):
    sh = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(rowh * len(rows)))
    tbl = sh.table
    for j, cw in enumerate(colw):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(rowh)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.text = ""
            p = c.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = str(val)
            r.font.size = Pt(size)
            r.font.name = "Calibri"
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            if j > 1 or (j == 1 and len(row) == 3):
                p.alignment = PP_ALIGN.CENTER
            c.fill.solid()
            if i == 0:
                c.fill.fore_color.rgb = NAVY
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.bold = True
            else:
                c.fill.fore_color.rgb = LIGHT if i % 2 else RGBColor(255, 255, 255)
                r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
                if winner_col is not None and j == winner_col:
                    r.font.bold = True
    return tbl


def pic(slide, name, x, y, w=None, h=None):
    slide.shapes.add_picture(
        f"{HERE}/{name}.png", Inches(x), Inches(y), width=Inches(w) if w else None, height=Inches(h) if h else None
    )


def foot(slide, text):
    tb(slide, 0.6, 6.95, 12.1, 0.4, text, size=10)


pct = lambda a, n: f"{a / n:.0%}"
a, d = S[CAND], S[BASE]

# 1 summary
MODE = {
    "all_on": "every chat flag on for every turn",
    "needed_only": "only the chat flags each case needs",
    "always_rag": "the flags each case needs, plus central RAG always on (the app default)",
}
mode = W[CAND]["flag_mode"]
s = new(f"{CL} vs {BL}: evaluation summary", f"Run {RUN}: {MODE.get(mode, mode)}")
cards = [
    (
        "Cost",
        "What each environment spends to answer the same questions: the answering model's calls plus background calls, including prompt caching.",
    ),
    (
        "Performance",
        "How quickly a user gets an answer: time to the first streamed text, and total time for the full reply.",
    ),
    (
        "Quality",
        "Whether answers match a written expected answer, judged blind by an LLM, plus which of the two answers is better.",
    ),
]
for i, (h, body) in enumerate(cards):
    x = 0.6 + i * 4.1
    box = s.shapes.add_shape(1, Inches(x), Inches(1.9), Inches(3.9), Inches(2.4))
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT
    box.line.fill.background()
    tb(s, x + 0.2, 2.0, 3.5, 0.5, h, size=22, color=NAVY, bold=True)
    tb(s, x + 0.2, 2.65, 3.5, 1.6, body, size=14)
tb(
    s,
    0.6,
    4.7,
    12.1,
    2,
    [
        [
            ("Scope: ", True),
            (
                f"{n_scen} scenarios, {n_cases} cases and {n_turn_defs} turns, run {reps} times each: {per_env_turns} answered turns per environment ({2 * per_env_turns} in all).",
                False,
            ),
        ],
        [
            ("Environments: ", True),
            (
                f"{CL} against {BL}: same cases, same users, run at the same time.",
                False,
            ),
        ],
        [("Scenarios: ", True), (", ".join(x.replace("_", " ") for x in scen) + ".", False)],
    ],
    size=15,
)

# 2 comparison
win = lambda x, y, lower=True: CAND if ((x < y) == lower and x != y) else (BASE if x != y else "tie")
rows = [["Dimension", "Measure", CL, BL, "Winner"]]


def add(dim, meas, x, y, fx, lower=True):
    w = win(x, y, lower)
    rows.append([dim, meas, fx(x), fx(y), LABEL[w]])


add("Cost", "Total cost (answers + other calls)", a["total"], d["total"], lambda v: f"${v:.2f}")
add("", "Cache-read share of input tokens", a["cache"], d["cache"], lambda v: f"{v:.0%}", lower=False)
add("Performance", "Time to first token, median", a["ttft50"], d["ttft50"], lambda v: f"{v:.1f}s")
add("", "Time to first token, p95", a["ttft95"], d["ttft95"], lambda v: f"{v:.1f}s")
add("", "Total time, median", a["tt50"], d["tt50"], lambda v: f"{v:.1f}s")
add("", "Total time, p95", a["tt95"], d["tt95"], lambda v: f"{v:.1f}s")
add(
    "Quality",
    "Answers judged acceptable",
    a["acc"] / a["acc_n"],
    d["acc"] / d["acc_n"],
    lambda v: f"{v:.0%}",
    lower=False,
)
rows.append(
    [
        "",
        "Pairwise: better answer (turns won)",
        f"{wins.get(CAND, 0)}",
        f"{wins.get(BASE, 0)}",
        LABEL[CAND if wins.get(CAND, 0) > wins.get(BASE, 0) else BASE] + f" ({wins.get('tie', 0)} ties)",
    ]
)
wins_by = pd.Series([r[4].split(" (")[0] for r in rows[1:]]).value_counts().to_dict()
s = new(
    "Comparison",
    f"{CL} wins {wins_by.get(LABEL[CAND], 0)} of {len(rows) - 1} measures, {BL} {wins_by.get(LABEL[BASE], 0)}",
)
table(s, 0.6, 1.8, 12.1, rows, [2.0, 4.6, 1.8, 1.8, 1.9], size=14, rowh=0.5, winner_col=4)
unlabelled = [LABEL[e] for e in (CAND, BASE) if not W[e].get("has_purpose_column", True)]
floor_note = (
    f"{' and '.join(unlabelled)}: background calls are not labelled by purpose and some (such as relevance checks and filter "
    "selection) may not be metered, so its total cost is a floor. "
    if unlabelled
    else ""
)
foot(
    s,
    floor_note + "Cache share is the share of all input tokens read from the prompt cache.",
)

# 3 cost
s = new("Cost", f"{CL} costs {a['total'] / d['total']:.0%} as much as {BL} over the same {per_env_turns} turns")
rows = [
    ["", CL, BL],
    ["Answer cost", f"${a['answer']:.2f}", f"${d['answer']:.2f}"],
    ["Other calls", f"${a['other']:.2f}", f"${d['other']:.2f}"],
    ["Total", f"${a['total']:.2f}", f"${d['total']:.2f}"],
    ["Cache-read share", f"{a['cache']:.0%}", f"{d['cache']:.0%}"],
]
table(s, 0.6, 1.9, 5.3, rows, [2.1, 1.6, 1.6], size=15, rowh=0.55)
tb(
    s,
    0.6,
    4.85,
    5.3,
    2.0,
    [
        "Answer cost is what the answering model calls cost. Other calls are background work: relevance checks, filter selection and similar.",
        f"Share of input read from the prompt cache: {CL} {a['cache']:.0%}, {BL} {d['cache']:.0%}.",
    ],
    size=13,
)
pic(s, "cum", 6.2, 1.7, w=6.6)
foot(
    s,
    "Cost = input price x (input + 0.1 x cache read + 1.25 x cache write tokens) + output price x output tokens."
    + (" Other calls for " + " and ".join(unlabelled) + " are an undercount." if unlabelled else ""),
)

# 4 performance
s = new("Performance", f"Median total time {a['tt50']:.1f}s for {CL} against {d['tt50']:.1f}s for {BL}")
rows = [
    ["Seconds", CL, BL],
    ["Time to first token, median", f"{a['ttft50']:.1f}", f"{d['ttft50']:.1f}"],
    ["Time to first token, p95", f"{a['ttft95']:.1f}", f"{d['ttft95']:.1f}"],
    ["Total time, median", f"{a['tt50']:.1f}", f"{d['tt50']:.1f}"],
    ["Total time, p95", f"{a['tt95']:.1f}", f"{d['tt95']:.1f}"],
]
table(s, 0.6, 1.5, 6.0, rows, [3.2, 1.4, 1.4], size=13, rowh=0.36)
r_al = W[CAND]
tb(
    s,
    6.9,
    1.6,
    5.9,
    2.0,
    [
        f"Wall-clock for the whole run: {CL} {(pd.Timestamp(W[CAND]['finished_at']) - pd.Timestamp(W[CAND]['started_at'])).total_seconds() / 60:.0f} min, {BL} {(pd.Timestamp(W[BASE]['finished_at']) - pd.Timestamp(W[BASE]['started_at'])).total_seconds() / 60:.0f} min, with several chats in flight at once.",
        "Timings are client-side; time to first token is the first non-empty streamed text. Delays between turns and uploads are excluded.",
    ],
    size=13,
)
pic(s, "ttft", 0.5, 3.5, w=5.9)
pic(s, "total", 6.8, 3.5, w=5.9)

# 5 quality
s = new(
    "Quality",
    f"{a['acc'] / a['acc_n']:.0%} of {CL} answers acceptable against {d['acc'] / d['acc_n']:.0%} for {BL}",
)
tb(
    s,
    0.6,
    1.5,
    12.1,
    1.3,
    [
        "Every turn has a written expected answer. A separate LLM judge ("
        + str(acc.judge.dropna().iloc[0])
        + ") saw only the question, the expected answer and the two responses, shuffled so it could not tell which environment wrote which. For each turn it decided whether each answer is acceptable and which is better, or a tie. {n_swaps} packets were repeated with the order swapped to check for position bias. The judge never saw earlier turns, tool calls or costs."
    ],
    size=13,
)
pic(s, "acc", 0.6, 2.7, w=5.6)
pic(s, "pair", 6.9, 2.7, w=5.6)
foot(
    s,
    f"Pairwise over {sum(wins.values())} turns: {CL} {wins.get(CAND, 0)}, {BL} {wins.get(BASE, 0)}, tie {wins.get('tie', 0)}. Bars show the share of turns in each scenario.",
)

# appendix: bias
s = new("Appendix: judge bias checks", "How far the judge can be trusted on this run")
rows = [
    ["Check", "Result"],
    ["Position: first answer (A) won", f"{pos.get('A', 0)}"],
    ["Position: second answer (B) won", f"{pos.get('B', 0)}"],
    ["Position: ties", f"{pos.get('tie', 0)}"],
    ["Swapped repeats of the same turn", f"{len(pairs)}"],
    ["Same verdict when A and B are swapped", f"{same} of {len(pairs)} ({same / len(pairs):.0%})"],
    ["Different verdicts that were tie vs a win", f"{tie_vs_win}"],
    ["Different verdicts that picked opposite winners", f"{flips}"],
]
table(s, 0.6, 1.8, 7.4, rows, [4.6, 2.8], size=14, rowh=0.5)
lean = (
    "no clear position bias"
    if abs(pos.get("A", 0) - pos.get("B", 0)) <= 0.1 * max(pos.get("A", 0) + pos.get("B", 0), 1)
    else "some position bias"
)
notes = [
    f"Judge verdicts across all packets: A won {pos.get('A', 0)}, B won {pos.get('B', 0)}, tie {pos.get('tie', 0)} ({lean})."
]
if flips == 0 and tie_vs_win:
    notes.append(
        f"All {tie_vs_win} inconsistent swaps were a tie one way and a win the other, never opposite winners. "
        f"{tw.count(CAND)} of them went to {CL} and {tw.count(BASE)} to {BL}."
    )
elif flips:
    notes.append(
        f"{flips} swaps gave opposite winners and {tie_vs_win} were a tie in one order and a win in the other."
    )
tb(s, 8.4, 1.8, 4.4, 4.5, notes, size=14)

# appendix: examples
EX = []
for spec in args.examples.split(";"):
    cid, rep, ti = spec.strip().split(":")
    sel = t[(t.case_id == cid) & (t.repeat == int(rep)) & (t.turn_index == int(ti))]
    ok = acc[(acc.case_id == cid) & (acc.repeat == int(rep)) & (acc.turn_index == int(ti))].set_index("env").acceptable
    pwin = pw[(pw.case_id == cid) & (pw.repeat == int(rep)) & (pw.turn_index == int(ti))].winner.iloc[0]
    okd = {e: bool(ok[e]) for e in (CAND, BASE)}
    tag = (
        "Both acceptable"
        if all(okd.values())
        else "Neither acceptable"
        if not any(okd.values())
        else f"{LABEL[CAND if okd[CAND] else BASE]} acceptable, {LABEL[BASE if okd[CAND] else CAND]} not"
    )
    EX.append((cid, int(rep), int(ti), f"{tag}; pairwise: {'tie' if pwin == 'tie' else LABEL[pwin] + ' wins'}"))
cl = lambda x, n: (x if len(x) <= n else x[: n - 1].rsplit(" ", 1)[0] + "...")
for cid, rep, ti, tag in EX:
    sel = t[(t.case_id == cid) & (t.repeat == rep) & (t.turn_index == ti)]
    g = lambda e: sel[sel.env == e].iloc[0]
    j = J[(J.case_id == cid) & (J.repeat == rep) & (J.turn_index == ti)]
    okv = {e: bool(j[(j.kind == "acceptable") & (j.env == e)].acceptable.iloc[0]) for e in (CAND, BASE)}
    reason = j[j.kind == "pairwise"].reason.iloc[0]
    winner = j[j.kind == "pairwise"].winner.iloc[0]
    s = new(f"Example: {cid}, turn {ti + 1}", tag)
    tb(
        s,
        0.6,
        1.5,
        12.1,
        1.2,
        [
            [("Question: ", True), (cl(g(CAND)["query"], 260), False)],
            [("Expected: ", True), (cl(g(CAND)["expected_answer"], 420), False)],
        ],
        size=12,
    )
    for k, e in enumerate((CAND, BASE)):
        x = 0.6 + k * 6.15
        tb(
            s,
            x,
            2.85,
            5.9,
            0.4,
            f"{LABEL[e]}: {'acceptable' if okv[e] else 'not acceptable'}",
            size=15,
            color=NAVY,
            bold=True,
        )
        tb(s, x, 3.3, 5.9, 2.7, cl(str(g(e)["response_text"]).replace("**", ""), 760), size=11)
    tb(
        s,
        0.6,
        6.0,
        12.1,
        1.2,
        [[("Judge: ", True), (f"{'tie' if winner == 'tie' else LABEL[winner] + ' wins'}. " + cl(reason, 420), False)]],
        size=12,
    )

prs.save(OUT + f"evaluation-summary-{RUN}.pptx")
print(f"Wrote {OUT}evaluation-summary-{RUN}.pptx")
